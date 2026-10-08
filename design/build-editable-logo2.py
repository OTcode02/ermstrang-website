#!/usr/bin/env python3
"""Bouwt het bewerkbare logo-bestand (PPTX) voor Canva en PowerPoint.

Waarom PPTX: Canva kan tekst in een geüpload SVG of PNG niet bewerken, zo'n bestand komt
binnen als plaatje. Een PPTX wordt bij import omgezet naar een ontwerp waarin de tekstvakken
echte, aanpasbare tekst blijven. Het merkteken gaat mee als afbeelding (dat hoort vast te blijven).
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
A = os.path.join(ROOT, "assets", "brand")

NAVY  = RGBColor(0x08,0x21,0x3F)
PAPER = RGBColor(0xF4,0xF7,0xFA)
CYAN  = RGBColor(0x00,0xAE,0xD0)
DIM   = RGBColor(0xCF,0xD9,0xE3)

PX = 100.0                                  # 100 px = 1 inch, dus maten zijn in "ontwerppixels"
def IN(px): return Inches(px / PX)
def PT(px): return Pt(px * 0.72)            # 1 px = 0,72 pt

def track(run, px_tracking):
    """Letter-spatiëring: python-pptx heeft er geen API voor, dus via het spc-attribuut (1/100 pt)."""
    run._r.get_or_add_rPr().set('spc', str(int(px_tracking * 0.72 * 100)))

def textbox(slide, x, y, w, h, lines, color, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines = [[(tekst, font, px, bold, tracking), ...], ...]  (binnenste lijst = één regel, meerdere runs)"""
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.05
        for txt, font, px, bold, tr in runs:
            r = p.add_run(); r.text = txt
            r.font.name, r.font.bold = font, bold
            r.font.size = PT(px); r.font.color.rgb = color
            track(r, tr)
    return tb

def blank(prs, bg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill; f.solid(); f.fore_color.rgb = bg
    return s

prs = Presentation()
prs.slide_width, prs.slide_height = IN(1584), IN(396)
W = lambda px=56, tr=-1.4: [("Ermstrang ", "Archivo", px, True, tr), ("Technologies", "Archivo", px, False, tr)]

# 1: LinkedIn-omslag (op navy)
s = blank(prs, NAVY)
s.shapes.add_picture(os.path.join(A,"mark-navy-512.png"), IN(64), IN(148), height=IN(100))
textbox(s, 200, 146, 620, 50, [W()], PAPER)
textbox(s, 800, 128, 748, 110, [[("Automatiseer het gedoe.", "Archivo", 46, True, -1.3)],
                                [("Focus op je vak.", "Archivo", 46, True, -1.3)]], PAPER)
textbox(s, 800, 240, 748, 26, [[("AI-assistenten op maat voor zzp en mkb · ermstrangtechnologies.nl",
                                "Source Sans 3", 19, False, 0)]], DIM)

# 2: liggend lockup op papier + uitleg dat de tekst aanpasbaar is
s = blank(prs, PAPER)
s.shapes.add_picture(os.path.join(A,"mark-paper-512.png"), IN(64), IN(148), height=IN(100))
textbox(s, 200, 146, 620, 50, [W()], NAVY)
textbox(s, 800, 140, 748, 120, [[("Pas de naam of de ondertitel aan in het tekstvak.", "Archivo", 28, True, -0.8)],
                                [("Het merkteken links is een afbeelding en blijft zoals het is.", "Source Sans 3", 19, False, 0)]], NAVY)

# 3: staand lockup (op navy)
s = blank(prs, NAVY)
s.shapes.add_picture(os.path.join(A,"mark-navy-512.png"), IN(762), IN(48), height=IN(150))
textbox(s, 0, 222, 1584, 56, [W(px=58)], PAPER, PP_ALIGN.CENTER)
textbox(s, 0, 292, 1584, 24, [[("AI-assistenten op maat voor zzp en mkb", "Source Sans 3", 19, False, 1.2)]], CYAN, PP_ALIGN.CENTER)

# 4: hoe je dit bestand gebruikt
s = blank(prs, PAPER)
textbox(s, 64, 60, 1450, 60, [[("Zo werkt het", "Archivo", 44, True, -1.2)]], NAVY)
textbox(s, 64, 140, 1450, 200, [
    [("In Canva: Bestand → Importeren. Canva zet de tekstvakken om naar echte Canva-tekst.", "Source Sans 3", 21, False, 0)],
    [("Lettertype Archivo staat in Canva bij de fonts (zoek op Archivo); Source Sans 3 ook.", "Source Sans 3", 21, False, 0)],
    [("In PowerPoint: dubbelklik op een tekstvak en typ. Het merkteken is een los plaatje.", "Source Sans 3", 21, False, 0)],
    [("In Inkscape/Figma/Illustrator: gebruik de SVG's in assets/img, daar is de tekst ook echt tekst.", "Source Sans 3", 21, False, 0)],
    [("Niet doen: het merkteken vervormen, van kleur veranderen of de puntjes los herschikken.", "Source Sans 3", 21, False, 0)],
], NAVY)

out = os.path.join(A, "ermstrang-logo-bewerkbaar.pptx")
prs.save(out)
print("geschreven:", out)
print("bytes:", os.path.getsize(out), "| slides:", len(prs.slides._sldIdLst))
